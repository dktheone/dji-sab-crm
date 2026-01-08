import os
from django.conf import settings
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.core.mail import send_mail
from django.template.loader import render_to_string
from django.contrib.auth.models import User, Group
from django.contrib.auth import authenticate, login, logout, get_user_model
from django.contrib.auth.views import PasswordResetView
from django.views.decorators.http import require_POST
from django.db.models import Sum, Q, F, Value, DecimalField, Max
from django.db.models import OuterRef, Subquery, Value, CharField
from django.db.models.functions import Coalesce
from django.db import models
from decimal import Decimal
from datetime import datetime, timedelta, date
import calendar

from apps.emp.forms import DepartmentForm, DesignationForm, SitesForm, CustomLoginForm, CustomPasswordResetForm, Contacts, ContactForm
from apps.emp.models import Department, Designation, Site
from apps.emp.models import Employee, EmployeeUpload, EmployeeStatus, EmployeeEducation, EmployeeExperience, EmployeeFamily, EmployeeSalaryMaster, EmployeeAdjustment, EmployeeSalaryTransaction, EmployeeAdjustmentMaster
from apps.emp.forms import EmployeeForm, EmployeeUploadForm, EmployeeUserForm, EmployeeEducationForm, EmployeeExperienceForm, EmployeeFamilyForm, EmployeeSalaryMasterForm, EmployeeAdjustmentForm, SalaryPreparationForm
from apps.emp.utils import role_required, fetch_ifsc_details
from country_state_city import City

def custom_login(request):
    # print('login request POST', request.POST, request.method)
    if request.method == 'POST':
        form = CustomLoginForm(data=request.POST)
        if form.is_valid():
            username = form.cleaned_data['username']
            password = form.cleaned_data['password']
            # print(username, password)
            user = authenticate(request, username=username, password=password)
            if user is not None:
                login(request, user)
                
                # ROLE-BASED LANDING PAGES
                user_role = user.groups.first().name if user.groups.exists() else None
                redirect_url = '/dashboard/'  # Default
                
                if user_role == 'hr':
                    # HR → Employee Dashboard
                    redirect_url = '/emp/dash/'
                
                elif user_role in ['employee', 'manager']:
                    # Employee/Manager → Their own activity page
                    try:
                        employee = Employee.objects.get(emp_code=user.username, status='Active')
                        redirect_url = f'/leads/activity/{employee.id}/'
                    except Employee.DoesNotExist:
                        # Fallback to leads list if no employee record
                        redirect_url = '/leads/'
                
                elif user_role == 'admin':
                    # Admin → Dashboard
                    redirect_url = '/dashboard/'
                
                return JsonResponse({'status': 'success', 'redirect': redirect_url})
            else:
                return JsonResponse({'status': 'error', 'message': 'Invalid employee code or password'}, status=400)
        else:
            return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
    else:
        form = CustomLoginForm()
    return render(request, 'login.html', {'form': form})



def custom_logout(request):
    logout(request)
    # messages.success(request, "You have been logged out successfully.")
    return redirect('/login')  # Name of your login URL
    # return render(request, 'logout.html')

        
# def dashboard(request):
#     return render(request, "emp/dashboard.html")

def dashboard(request):
    """Employee dashboard view - includes upcoming follow-ups and role-based stats"""
    from datetime import date, timedelta
    from django.db.models import Count, Q
    from apps.leads.models import Lead, LeadStatus, FollowUp
    from apps.emp.models import Employee, Site
    from apps.vendors.models import Vendor, VendorStatus
    
    employee_id = request.user
    
    if request.user.is_authenticated:
        employee = get_object_or_404(Employee, emp_code=request.user.username) if isinstance(request.user, User) and not request.user.is_superuser else None
        # Handle superuser falling back to an employee profile if linked, or just proceed
        if not employee and not request.user.is_superuser:
             # Try to find employee by email if username match fails, or handle gracefully
             pass

        # Get user role
        user_groups = list(request.user.groups.values_list('name', flat=True))
        is_admin = 'admin' in user_groups or request.user.is_superuser
        is_hr = 'hr' in user_groups
        is_manager = 'manager' in user_groups
        
        # --- Common Context: Upcoming Follow-ups (Personal) ---
        today = date.today()
        next_week = today + timedelta(days=7)
        
        if is_admin or is_hr or is_manager:
            # Show ALL records
            upcoming_leads = Lead.objects.filter(
                next_follow_up__gte=today,
                lead_status__in=[LeadStatus.NEW, LeadStatus.IN_PROGRESS, LeadStatus.QUOTED]
            ).select_related('created_by').order_by('next_follow_up')

            recent_followups = FollowUp.objects.all().select_related('lead', 'created_by').order_by('-follow_up_date', '-created_at')[:15]
            
            # For admin/hr visuals, we still want to separate overdue
            overdue_leads = Lead.objects.filter(
                next_follow_up__lt=today,
                lead_status__in=[LeadStatus.NEW, LeadStatus.IN_PROGRESS, LeadStatus.QUOTED]
            ).order_by('next_follow_up')
        else:
            # Regular employees only see their own
            if employee:
                upcoming_leads = Lead.objects.filter(
                    created_by=employee,
                    next_follow_up__gte=today,
                    lead_status__in=[LeadStatus.NEW, LeadStatus.IN_PROGRESS, LeadStatus.QUOTED]
                ).select_related('created_by').order_by('next_follow_up')
                
                recent_followups = FollowUp.objects.filter(created_by=employee).select_related('lead', 'created_by').order_by('-follow_up_date', '-created_at')[:15]
                
                overdue_leads = Lead.objects.filter(
                    created_by=employee,
                    next_follow_up__lt=today,
                    lead_status__in=[LeadStatus.NEW, LeadStatus.IN_PROGRESS, LeadStatus.QUOTED]
                ).order_by('next_follow_up')
            else:
                upcoming_leads = Lead.objects.none()
                recent_followups = FollowUp.objects.none()
                overdue_leads = Lead.objects.none()

        # Separate lists for clearer UI (Optional refinement of upcoming vs today can be done in template or here)
        # For the Grid requirement, we pass the main QuerySets
        
        # --- Personal Financial Details (For all users if employee linked) ---
        current_salary = None
        active_adjustments = None
        
        if employee:
            from apps.emp.models import EmployeeSalaryMaster, EmployeeAdjustmentMaster
            current_salary = EmployeeSalaryMaster.objects.filter(employee=employee, status='Active').first()
            active_adjustments = EmployeeAdjustmentMaster.objects.filter(
                employee=employee, 
                status='Active'
            ).order_by('-date_issued')

        context = {
            'employee': employee,
            'is_admin': is_admin,
            'is_hr': is_hr,
            'is_manager': is_manager,
            'view_name': request.resolver_match.view_name,
            'overdue_leads': overdue_leads,
            'upcoming_leads': upcoming_leads, # Now includes Today+
            'recent_followups': recent_followups,
            'total_my_tasks': overdue_leads.count() + upcoming_leads.count(),
            'current_salary': current_salary,
            'active_adjustments': active_adjustments,
        }

        # --- Admin/HR Context: Overall Statistics ---
        if is_admin or is_hr:
            # Employee Stats
            total_employees = Employee.objects.count()
            active_employees = Employee.objects.filter(status='Active').count()
            
            # Site Stats
            total_sites = Site.objects.count()
            active_sites = Site.objects.filter(status='Active').count()
            
            # Vendor Stats
            total_vendors = Vendor.objects.count()
            active_vendors = Vendor.objects.filter(status=VendorStatus.ACTIVE).count()
            
            # Lead Stats (Aggregated)
            lead_stats = Lead.objects.aggregate(
                total=Count('lead_id'),
                new=Count('lead_id', filter=Q(lead_status=LeadStatus.NEW)),
                converted=Count('lead_id', filter=Q(lead_status=LeadStatus.CONVERTED)),
            )

            context.update({
                'stats_total_employees': total_employees,
                'stats_active_employees': active_employees,
                'stats_total_sites': total_sites,
                'stats_active_sites': active_sites,
                'stats_total_vendors': total_vendors,
                'stats_active_vendors': active_vendors,
                'stats_leads_total': lead_stats['total'],
                'stats_leads_new': lead_stats['new'],
                'stats_leads_converted': lead_stats['converted'],
            })
        
        return render(request, 'emp/dashboard.html', context)
    else:
        return custom_logout(request)
    

# Password Change View (for logged-in users and admin changing user passwords)
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import PasswordChangeForm, SetPasswordForm
from django.contrib.auth import update_session_auth_hash

@login_required
def password_change(request, employee_id=None):
    """
    Handle password change for:
    1. Users changing their own password (employee_id=None or their own ID)
    2. Admin/HR changing another user's password (employee_id provided)
    """
    # Get current user's role
    user_role = request.user.groups.first().name if request.user.groups.exists() else None
    
    # Determine target employee and user
    if employee_id:
        # Someone is trying to change a specific employee's password
        try:
            target_employee = Employee.objects.get(id=employee_id)
            target_user = User.objects.get(username=target_employee.emp_code)
        except (Employee.DoesNotExist, User.DoesNotExist):
            messages.error(request, 'Employee or user not found')
            return redirect('/dashboard/')
        
        # Check permissions
        if user_role not in ['admin', 'hr']:
            # Non-admin/hr users can only change their own password
            current_employee = Employee.objects.filter(emp_code=request.user.username).first()
            if not current_employee or current_employee.id != employee_id:
                messages.error(request, 'You can only change your own password')
                return redirect('/dashboard/')
    else:
        # No employee_id provided - user changing their own password
        try:
            target_employee = Employee.objects.get(emp_code=request.user.username)
            target_user = request.user
        except Employee.DoesNotExist:
            # User might be admin without employee record
            target_employee = None
            target_user = request.user
    
    # Check if user is changing their own password
    is_self = target_user == request.user
    
    if request.method == 'POST':
        if is_self:
            # User changing own password - requires old password
            form = PasswordChangeForm(request.user, request.POST)
        else:
            # Admin changing user password - no old password required
            form = SetPasswordForm(target_user, request.POST)
        
        if form.is_valid():
            user = form.save()
            
            # Keep current user logged in if they changed their own password
            if is_self:
                update_session_auth_hash(request, user)
            
            message = 'Password changed successfully' if is_self else f'Password updated for {target_user.get_full_name() or target_user.username}'
            return JsonResponse({'status': 'success', 'message': message})
        else:
            return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
    else:
        if is_self:
            form = PasswordChangeForm(request.user)
        else:
            form = SetPasswordForm(target_user)
    
    return render(request, 'password_change.html', {
        'form': form,
        'target_user': target_user,
        'target_employee': target_employee,
        'is_self': is_self
    })

# Password Reset Views (for forgot password flow)
from django.contrib.auth.views import (
    PasswordResetView, 
    PasswordResetDoneView,
    PasswordResetConfirmView,
    PasswordResetCompleteView
)
from django.urls import reverse_lazy

class CustomPasswordResetView(PasswordResetView):
    template_name = 'password_reset_form.html'
    email_template_name = 'emp/emails/password_reset.html' # Use HTML file for text body (fallback) - ideally should be .txt
    html_email_template_name = 'emp/emails/password_reset.html' # Use HTML file for HTML body
    subject_template_name = 'email/password_reset_subject.txt'
    success_url = reverse_lazy('password_reset_done')
    
class CustomPasswordResetDoneView(PasswordResetDoneView):
    template_name = 'password_reset_done.html'
    
class CustomPasswordResetConfirmView(PasswordResetConfirmView):
    template_name = 'password_reset_confirm.html'
    success_url = reverse_lazy('password_reset_complete')
    
class CustomPasswordResetCompleteView(PasswordResetCompleteView):
    template_name = 'password_reset_complete.html'


@role_required('admin', 'hr')
def department_crud(request):
    if request.method == 'POST':
        if request.POST.get('edit_id'):
            # Handle edit form submission
            department = get_object_or_404(Department, id=request.POST['edit_id'])
            form = DepartmentForm(request.POST, instance=department)
            if form.is_valid():
                department = form.save(commit=False)
                department.created_by = request.user
                department.save()
                return JsonResponse({'status': 'success', 'message': 'Department updated successfully'})
            else:
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
        else:
            # Handle add form submission
            form = DepartmentForm(request.POST)
            if form.is_valid():
                department = form.save(commit=False)
                department.created_by = request.user
                department.save()
                return JsonResponse({'status': 'success', 'message': 'Department added successfully'})
            else:
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
    else:
        add_form = DepartmentForm()
        edit_form = DepartmentForm()  # Empty form for edit modal
        departments = Department.objects.all()
        return render(request, 'emp/department_crud.html', {
            'add_form': add_form,
            'edit_form': edit_form,
            'departments': departments,
            'view_name': request.resolver_match.view_name
        })

@role_required('admin')
def delete_department(request, department_id):
    department = get_object_or_404(Department, id=department_id)
    if request.method == 'POST':
        try:
            department.delete()
            return JsonResponse({'status': 'success', 'message': 'Department deleted successfully'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
    return JsonResponse({'status': 'error', 'message': 'Invalid request method'}, status=400)

@role_required('admin', 'hr')
def designation_crud(request):
    if request.method == 'POST':
        if request.POST.get('edit_id'):
            # Handle edit form submission
            designation = get_object_or_404(Designation, id=request.POST['edit_id'])
            print('designation', designation)
            print('request.POST', request.POST)
            form = DesignationForm(request.POST, instance=designation)
            print('form', form, form.is_valid())
            if form.is_valid():
                designation = form.save(commit=False)
                designation.created_by = request.user
                designation.save()
                messages.success(request, message="Designation updated successfully")
                # return render(request, 'emp/designation_crud.html')
                return JsonResponse({'status': 'success', 'message': 'Designation updated successfully'})
            else:
                messages.error(message="Designation Updation Failed!")
                # return render(request, 'emp/designation_crud.html')
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
        else:
            # Handle add form submission
            form = DesignationForm(request.POST)
            if form.is_valid():
                designation = form.save(commit=False)
                designation.created_by = request.user
                designation.save()
                messages.success(request, message="Designation Added successfully")
                # return render(request, 'emp/designation_crud.html')
                return JsonResponse({'status': 'success', 'message': 'Designation added successfully'})
            else:
                messages.error(request, message="Designation Insertion Failed!")
                # return render(request, 'emp/designation_crud.html')
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
    else:
        add_form = DesignationForm()
        edit_form = DesignationForm()  # Empty form for edit modal
        designations = Designation.objects.all()
        return render(request, 'emp/designation_crud.html', {
            'add_form': add_form,
            'edit_form': edit_form,
            'designations': designations,
            'view_name': request.resolver_match.view_name
        })

@role_required('admin')
def delete_designation(request, designation_id):
    designation = get_object_or_404(Designation, id=designation_id)
    if request.method == 'POST':
        try:
            designation.delete()
            return JsonResponse({'status': 'success', 'message': 'Designation deleted successfully'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
    return JsonResponse({'status': 'error', 'message': 'Invalid request method'}, status=400)

@role_required('admin', 'hr', 'employee')
def site_crud(request, site_id=None):
    site_details = get_object_or_404(Site, id=site_id) if site_id else None
    if request.method == 'POST':
        if request.POST.get('edit_id'):
            # Handle edit form submission
            site = get_object_or_404(Site, id=request.POST['edit_id'])
            form = SitesForm(request.POST, instance=site)
            if form.is_valid():
                site = form.save(commit=False)
                site.created_by = request.user
                site.save()
                messages.success(request, "Site updated successfully")
                return JsonResponse({'status': 'success', 'message': 'Site updated successfully'})
            else:
                messages.error(request, "Site Updation Failed!")
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
        else:
            # Handle add form submission
            form = SitesForm(request.POST)
            if form.is_valid():
                site = form.save(commit=False)
                site.created_by = request.user
                site.save()
                messages.success(request, "New Site Added successfully")
                return JsonResponse({'status': 'success', 'message': 'New Site added successfully'})
            else:
                messages.error(request, "New Site Insertion Failed!")
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
    else:
        add_form = SitesForm()
        edit_form = SitesForm(instance=site_details) if site_details else SitesForm()
        sites = Site.objects.all()
        initial_city = site_details.city if site_details else None
        initial_state = site_details.state if site_details else None
        return render(request, 'emp/sites_crud.html', {
            'site_add_form': add_form,
            'site_edit_form': edit_form,
            'sites': sites,
            'initial_city': initial_city,
            'initial_state': initial_state,
            'site_details': site_details,
            'view_name': request.resolver_match.view_name,
        })

@role_required('admin', 'hr', 'employee')
def contacts_view(request):
    contacts = Contacts.objects.order_by('-created_at')
    total_count = contacts.count()

    # Calculate upcoming birthdays (next 7 days, dob or event_dob)
    today = datetime.today()
    upcoming_birthdays = []
    for i in range(1, 8):
        check_date = today + timedelta(days=i)
        contacts_with_bd = Contacts.objects.filter(
            Q(dob__month=check_date.month, dob__day=check_date.day) |
            Q(event_date__month=check_date.month, event_date__day=check_date.day)
        )
        for c in contacts_with_bd:
            bd_field = 'Birthday' if (c.dob and c.dob.month == check_date.month and c.dob.day == check_date.day) else c.event_date_remark or 'Event'
            upcoming_birthdays.append({
                'name': f"{c.first_name} {c.last_name}",
                'birthday': check_date.strftime('%m/%d'),
                'type': bd_field
            })

    context = {
        'contacts': contacts,
        'total_count': total_count,
        'upcoming_birthdays': upcoming_birthdays,
        'contact_form': ContactForm(),
        'view_name': request.resolver_match.view_name,
    }
    return render(request, 'emp/contacts.html', context)

@role_required('admin', 'hr', 'employee')
def contact_add(request):
    if request.method == 'POST':
        form = ContactForm(request.POST)
        if form.is_valid():
            contact = form.save(commit=False)
            contact.created_by = request.user
            contact.save()
            return JsonResponse({'success': True})
        else:
            return JsonResponse({'success': False, 'errors': form.errors})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@role_required('admin', 'hr', 'employee')
def get_site_details(request, site_id):
    site = get_object_or_404(Site, id=site_id)
    data = {
        'site_name': site.site_name,
        'full_address': f"{site.address or ''}, {site.city or ''}, {site.state or ''}, {site.pincode or ''}".strip(', '),
        'contact1': site.contact1 or 'N/A',
        'contact2': site.contact2 or 'N/A',
        'email': site.email or 'N/A',
        'website': site.website or 'N/A',
        'status': site.status,
        'remark': site.remark or 'No remarks',
    }
    return JsonResponse(data)




@role_required('admin')
def delete_site(request, site_id):
    site = get_object_or_404(Site, id=site_id)
    if request.method == 'POST':
        try:
            site.delete()
            return JsonResponse({'status': 'success', 'message': 'Site deleted successfully'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
    return JsonResponse({'status': 'error', 'message': 'Invalid request method'}, status=400)

@role_required('admin', 'hr', 'employee')
def employee_form(request, employee_id=None):
    employee = get_object_or_404(Employee, id=employee_id) if employee_id else None
    if request.method == 'POST':
        form = EmployeeForm(request.POST, request.FILES, instance=employee)
        if form.is_valid():
            employee = form.save(commit=False)
            employee.created_by = request.user
            employee.save()
            return JsonResponse({
                'status': 'success',
                'message': 'Employee saved successfully',
                'redirect': f'/emp/details/{employee.id}'
            })
        else:
            return JsonResponse({'status': 'error', 'errors': form.errors, 'test': '7275'}, status=400)
    else:
        employee_form = EmployeeForm(instance=employee)
        initial_city = employee.cor_city if employee else None
        initial_state = employee.cor_state if employee else None
        return render(request, 'emp/employee_form.html', {
            'employee_form': employee_form,
            'employee': employee,
            'initial_city': initial_city,
            'initial_state': initial_state,
            'view_name': request.resolver_match.view_name,
            'step': 1  # For wizard indicator
        })
# New view for additional details (Education, Experience, Family)
@role_required('admin', 'hr', 'employee')
def employee_additional(request, employee_id):
    employee = get_object_or_404(Employee, id=employee_id)
    # if request.user.groups.filter(name='employee').exists() and employee.created_by != request.user:
    #     return redirect('emp:employee_list')

    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'add_education' or action == 'edit_education':
            instance = get_object_or_404(EmployeeEducation, id=request.POST.get('id')) if action == 'edit_education' else None
            form = EmployeeEducationForm(request.POST, instance=instance)
            if form.is_valid():
                obj = form.save(commit=False)
                obj.employee = employee
                obj.created_by = request.user
                obj.save()
                return JsonResponse({'status': 'success', 'message': 'Education saved successfully'})
            else:
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
        elif action == 'delete_education':
            instance = get_object_or_404(EmployeeEducation, id=request.POST.get('id'))
            instance.delete()
            return JsonResponse({'status': 'success', 'message': 'Education deleted successfully'})
        # Similar for experience
        elif action == 'add_experience' or action == 'edit_experience':
            instance = get_object_or_404(EmployeeExperience, id=request.POST.get('id')) if action == 'edit_experience' else None
            form = EmployeeExperienceForm(request.POST, instance=instance)
            if form.is_valid():
                obj = form.save(commit=False)
                obj.employee = employee
                obj.created_by = request.user
                obj.save()
                return JsonResponse({'status': 'success', 'message': 'Experience saved successfully'})
            else:
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
        elif action == 'delete_experience':
            instance = get_object_or_404(EmployeeExperience, id=request.POST.get('id'))
            instance.delete()
            return JsonResponse({'status': 'success', 'message': 'Experience deleted successfully'})
        # Similar for family
        elif action == 'add_family' or action == 'edit_family':
            instance = get_object_or_404(EmployeeFamily, id=request.POST.get('id')) if action == 'edit_family' else None
            form = EmployeeFamilyForm(request.POST, instance=instance)
            if form.is_valid():
                obj = form.save(commit=False)
                obj.employee = employee
                obj.created_by = request.user
                obj.save()
                return JsonResponse({'status': 'success', 'message': 'Family member saved successfully'})
            else:
                return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
        elif action == 'delete_family':
            instance = get_object_or_404(EmployeeFamily, id=request.POST.get('id'))
            instance.delete()
            return JsonResponse({'status': 'success', 'message': 'Family member deleted successfully'})

    education_form = EmployeeEducationForm()
    experience_form = EmployeeExperienceForm()
    family_form = EmployeeFamilyForm()
    educations = EmployeeEducation.objects.filter(employee=employee)
    experiences = EmployeeExperience.objects.filter(employee=employee)
    families = EmployeeFamily.objects.filter(employee=employee)

    return render(request, 'emp/employee_additional.html', {
        'employee': employee,
        'education_form': education_form,
        'experience_form': experience_form,
        'family_form': family_form,
        'educations': educations,
        'experiences': experiences,
        'families': families,
        'view_name': request.resolver_match.view_name,
        'step': 2  # For wizard indicator
    })
    
    
@role_required('admin', 'hr', 'employee')
def employee_uploads(request, employee_id):
    employee = get_object_or_404(Employee, id=employee_id)
    # if request.user.groups.filter(name='employee').exists() and employee.created_by != request.user:
    #     return redirect('emp:employee_list')
    
    if request.method == 'POST':
        upload_form = EmployeeUploadForm(request.POST, request.FILES)
        if upload_form.is_valid():
            upload = upload_form.save(commit=False)
            upload.employee = employee
            upload.uploaded_by = request.user
            upload.created_by = request.user
            upload.save()
            return JsonResponse({'status': 'success', 'message': 'Document uploaded successfully'})
        else:
            return JsonResponse({'status': 'error', 'errors': upload_form.errors}, status=400)
    else:
        upload_form = EmployeeUploadForm()
        uploads = EmployeeUpload.objects.filter(employee=employee)
        return render(request, 'emp/employee_uploads.html', {
            'upload_form': upload_form,
            'employee': employee,
            'uploads': uploads,
            'view_name': request.resolver_match.view_name,
            'step': 3  # For wizard indicator
        })


@require_POST
def get_ifsc_data(request):
    ifsc_code = request.POST.get('ifsc_code', '').strip().upper() # Get from POST data, clean, and make uppercase

    ifsc_details = fetch_ifsc_details(ifsc_code)

    if "error" in ifsc_details:
        return JsonResponse({"status": "error", "message": ifsc_details["error"]}, status=400)
    else:
        return JsonResponse({"status": "success", "data": ifsc_details})


@role_required('admin', 'hr')
def employee_list(request):
    employees = Employee.objects.all()
    return render(request, 'emp/employee_list.html', {
        'employees': employees,
        'view_name': request.resolver_match.view_name,
        'status_choices': EmployeeStatus.choices
    })
    
    
# @role_required('admin', 'hr')
# def employee_view(request, employee_id):
#     employee = get_object_or_404(Employee, id=employee_id)
#     uploads = EmployeeUpload.objects.filter(employee=employee)
#     return render(request, 'emp/employee_view.html', {'employee': employee, 'uploads': uploads})

@role_required('admin', 'hr', 'employee')
def employee_view(request, employee_id):
    employee = get_object_or_404(Employee, id=employee_id)
    # if request.user.groups.filter(name='employee').exists() and employee.created_by != request.user:
    #     return redirect('emp:employee_list')
    
    educations = EmployeeEducation.objects.filter(employee=employee)
    experiences = EmployeeExperience.objects.filter(employee=employee)
    families = EmployeeFamily.objects.filter(employee=employee)
    uploads = EmployeeUpload.objects.filter(employee=employee)
    
    return render(request, 'emp/employee_view.html', {
        'employee': employee,
        'educations': educations,
        'experiences': experiences,
        'families': families,
        'uploads': uploads,
        'view_name': request.resolver_match.view_name,
        'step': 4  # For wizard indicator
    })


@role_required('admin', 'hr')
def update_employee_status(request, employee_id):
    print(request, employee_id, request.POST.get('status'))
    if request.method == 'POST':
        employee = get_object_or_404(Employee, id=employee_id)
        
        print(employee)
        
        status = request.POST.get('status')
        status_split =  list(status)
        print(status_split)
        print((status in [choice[0] for choice in EmployeeStatus.choices]), status)
        print([choice[0] for choice in EmployeeStatus.choices])
        
        # print(type(status), [type(choice[0]) for choice in EmployeeStatus.choices], (status in [choice[0] for choice in EmployeeStatus.choices]))
        # for choice in EmployeeStatus.choices:
        #     print(choice[0], status, (status == choice[0]) )
        #     print(len(choice[0]), len(status), (len(status) == len(choice[0])))
        
        if status in [choice[0] for choice in EmployeeStatus.choices]:
            print('status', status)
            employee.status = status
            employee.save()
            return JsonResponse({'status': 'success', 'message': 'Status updated successfully'})
        return JsonResponse({'status': 'error', 'message': 'Invalid status'}, status=400)
    return JsonResponse({'status': 'error', 'message': 'Invalid request method'}, status=400)





def get_cities(request):
    if request.method == 'GET':
        state_code = request.GET.get('state')  # Fixed: changed from 'state_code' to 'state'
        
        if state_code:
            cities = City.get_cities_of_state('IN', state_code)
            city_list = [city.name for city in cities]  # Return just the city names
            return JsonResponse({'cities': city_list})
        return JsonResponse({'cities': []}, status=400)
    return JsonResponse({'error': 'Invalid request method'}, status=400)



@role_required('admin', 'hr')
def create_employee_user(request):
    
    # User = get_user_model()
    if request.method == 'POST':
        form = EmployeeUserForm(request.POST)
        if form.is_valid():
            employee = form.cleaned_data['employee']
            try:
                user = User.objects.create_user(
                    username=employee.emp_code,
                    password=employee.emp_code,
                    email=employee.email or '',
                    first_name=employee.first_name,
                    last_name=employee.last_name
                )
                
                # Update emp_id using Raw SQL (Field not in Django Model)
                from django.db import connection
                with connection.cursor() as cursor:
                    cursor.execute("UPDATE auth_user SET emp_id = %s WHERE id = %s", [employee.id, user.id])
                
                # Assign Role based on Designation
                role_name = 'employee'
                if employee.designation:
                    d_name = employee.designation.designation_name.lower()
                    if 'hr' in d_name or 'human resource' in d_name:
                        role_name = 'hr'
                    elif 'manager' in d_name:
                        role_name = 'manager'
                
                try:
                    group = Group.objects.get(name=role_name)
                    user.groups.add(group)
                except Group.DoesNotExist:
                    # Fallback to employee
                    try:
                        group = Group.objects.get(name='employee')
                        user.groups.add(group)
                    except Group.DoesNotExist:
                        pass
                user.save()

                # Send Welcome Email
                if employee.email:
                    try:
                        subject = 'Welcome to SAB Hospitality - Your Login Details'
                        html_message = render_to_string('emp/emails/account_created.html', {
                            'employee_name': f"{employee.first_name} {employee.last_name}",
                            'username': user.username,
                            'password': user.username, # Password is same as emp_code initially
                        })
                        send_mail(
                            subject,
                            message='', # Plain text fallback (could be better)
                            from_email=None, # Uses DEFAULT_FROM_EMAIL
                            recipient_list=[employee.email],
                            html_message=html_message,
                            fail_silently=True
                        )
                    except Exception as e:
                        print(f"Error sending email: {e}")

                return JsonResponse({'status': 'success', 'message': f'User created for {employee.emp_code}'})
            except Exception as e:
                return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
        else:
            return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
    else:
        form = EmployeeUserForm()
        return render(request, 'emp/create_employee_user.html', {'form': form, 'view_name': request.resolver_match.view_name})
    
    

from django.views.generic import ListView
from django.db.models import Q, Case, When, Value, CharField, Count
from django.utils import timezone
    
# @role_required('admin', 'hr')
class EmployeeDashboardView(ListView):
    model = Employee
    template_name = 'emp/employee_dashboard.html'
    context_object_name = 'employees'
    paginate_by = 25

    def get_queryset(self):
        from dateutil.relativedelta import relativedelta
        
        queryset = super().get_queryset().filter(status__in=['Active', 'Onboarding', 'Probation'])
        
        # Check for PREVIOUS month payroll (since current month salary is prepared in next month)
        today = timezone.now()
        prev_month = today - relativedelta(months=1)
        check_month = prev_month.month
        check_year = prev_month.year

        from django.db.models import Exists, OuterRef
        
        # Annotate statuses using Subqueries to avoid Cartesian product (duplication)
        has_processed_transaction = Exists(
            EmployeeSalaryTransaction.objects.filter(
                employee=OuterRef('pk'),
                month=check_month,
                year=check_year
            )
        )
        
        has_active_salary_master = Exists(
            EmployeeSalaryMaster.objects.filter(
                employee=OuterRef('pk'),
                status='Active',
                effective_from__lte=timezone.now()
            )
        )
        
        has_pending_adjustment = Exists(
            EmployeeAdjustment.objects.filter(
                employee=OuterRef('pk'),
                status='Pending'
            )
        )

        queryset = queryset.annotate(
            payroll_status=Case(
                When(has_processed_transaction, then=Value('Processed')),
                default=Value('Pending'),
                output_field=CharField()
            ),
            login_status=Case(
                When(emp_code__in=User.objects.filter(is_active=True).values('username'), then=Value('Active')),
                When(emp_code__in=User.objects.filter(is_active=False).values('username'), then=Value('Disabled')),
                default=Value('Not Created'),
                output_field=CharField()
            ),
            salary_status=Case(
                When(has_active_salary_master, then=Value('Assigned')),
                default=Value('Pending'),
                output_field=CharField()
            ),
            adjustments_status=Case(
                When(has_pending_adjustment, then=Value('Pending')),
                default=Value('No Pending'),
                output_field=CharField()
            )
        ) # distinct() is no longer strictly necessary if joins are removed, but keeping it is safe.
        queryset = queryset.distinct()

        # Apply quick filters
        filter_type = self.request.GET.get('filter')
        if filter_type == 'payroll_pending':
            queryset = queryset.filter(payroll_status='Pending')
        elif filter_type == 'login_not_created':
            queryset = queryset.filter(login_status='Not Created')
        elif filter_type == 'salary_pending':
            queryset = queryset.filter(salary_status='Pending')
        elif filter_type == 'adjustments_pending':
            queryset = queryset.filter(adjustments_status='Pending')

        # Apply advanced filters
        dept = self.request.GET.get('dept')
        if dept:
            queryset = queryset.filter(department__department_name=dept)
        status = self.request.GET.get('status')
        if status:
            queryset = queryset.filter(status=status)

        # Search
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(first_name__icontains=search) |
                Q(last_name__icontains=search) |
                Q(emp_code__icontains=search) |
                Q(department__department_name__icontains=search)
            )

        return queryset.order_by('first_name')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        current_month = timezone.now().month
        current_year = timezone.now().year
        qs = Employee.objects.filter(status__in=['Active', 'Onboarding', 'Probation'])

        # Stat cards data
        context['stats'] = {
            'total_employees': qs.count(),
            'payroll_pending': qs.filter(
                ~Q(salary_transactions__month=current_month, salary_transactions__year=current_year)
            ).count(),
            'login_not_created': qs.filter(
                ~Q(emp_code__in=User.objects.values('username'))
            ).count(),
            'salary_processed': qs.filter(
                salary_masters__status='Active',
                salary_masters__effective_from__lte=timezone.now()
            ).count(),
            'adjustments_pending': qs.filter(adjustments__status='PENDING').count(),
            'view_name': '',
        }
        context['departments'] = Department.objects.filter(status='Active').values_list('department_name', flat=True)
        context['statuses'] = [choice[0] for choice in EmployeeStatus.choices]
        context['current_filter'] = self.request.GET.get('filter', 'all')
        return context   
    
    

# region SALARY_OLD
# @role_required('admin', 'hr')
# def salary_master_list(request):
#     employees = Employee.objects.filter(status='Active')
#     return render(request, 'emp/salary_master_list.html', {'employees': employees})

@role_required('admin', 'hr')
def salary_master_form(request, employee_id):
    # REDIRECT to the unified Detail/History view
    return redirect('apps.emp:salary_master_detail', emp_id=employee_id)

@role_required('admin', 'hr')
def adjustments_list(request):
    employees = Employee.objects.filter(status='Active')
    return render(request, 'emp/adjustments_list.html', {'employees': employees, 'view_name': request.resolver_match.view_name})

@role_required('admin', 'hr')
def adjustments_form(request, employee_id):
    from apps.emp.models import EmployeeAdjustmentMaster, EmployeeAdjustmentTransaction
    
    employee = get_object_or_404(Employee, id=employee_id)
    if request.method == 'POST':
        # Manually create EmployeeAdjustmentMaster (no form needed for now)
        try:
            adjustment_type = request.POST.get('adjustment_type')
            total_amount = Decimal(request.POST.get('amount', 0))
            date_issued = request.POST.get('date')
            description = request.POST.get('remarks', '')
            
            # Create master record
            EmployeeAdjustmentMaster.objects.create(
                employee=employee,
                adjustment_type=adjustment_type,
                total_amount=total_amount,
                date_issued=date_issued,
                description=description,
                status='Active',
                created_by=request.user
            )
            return JsonResponse({'status': 'success', 'message': 'Adjustment saved'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
    
    
    # Unified History Logic
    masters = EmployeeAdjustmentMaster.objects.filter(employee=employee)
    transactions = EmployeeAdjustmentTransaction.objects.filter(adjustment_master__employee=employee)
    
    history = []
    
    # Add Masters to history
    for m in masters:
        history.append({
            'date': m.date_issued,
            'type': m.adjustment_type,
            'amount': m.total_amount,
            'category': 'Master',
            'status': m.status,
            'description': m.description,
            'raw_obj': m
        })
        
    # Add Transactions to history
    for t in transactions:
        history.append({
            'date': t.settlement_date,
            'type': f"{t.transaction_type} Settlement",
            'amount': t.settlement_amount,
            'category': 'Transaction',
            'status': 'Completed',
            'description': f"{t.remarks} ({t.salary_month}/{t.salary_year})",
            'raw_obj': t
        })
    
    # Sort by date descending
    history.sort(key=lambda x: x['date'], reverse=True)
    
    # Combine for display (you can update template to show both)
    form = EmployeeAdjustmentForm()  # Keep form for now for template compatibility
    return render(request, 'emp/adjustments_form.html', {
        'form': form, 
        'employee': employee, 
        # 'adjustments': old_adjustments,  # Removed old
        'history': history,  # New unified history
        'view_name': request.resolver_match.view_name
    })


# @role_required('admin', 'hr')
# def salary_preparation(request):
#     month = request.GET.get('month', datetime.now().month)
#     year = request.GET.get('year', datetime.now().year)
#     site = request.GET.get('site', 'ALL')
#     months = [datetime(1900, i, 1).strftime('%B') for i in range(1, 13)]

#     employees = Employee.objects.filter(status='Active')
#     if site != 'ALL':
#         employees = employees.filter(site=site)

#     # Proper correlated subquery
#     salary_txn = EmployeeSalaryTransaction.objects.filter(
#         employee_id=OuterRef('id'),
#         month=month,
#         year=year
#     ).values('status')[:1]

#     employees = employees.annotate(
#         prepared=Coalesce(
#             Subquery(salary_txn, output_field=CharField()),
#             Value('Not Prepared'),
#             output_field=CharField()
#         )
#     )
#     # print(employees)
#     # for emp in employees:
#     #     print(emp, emp.id, emp.status)
#     sites = Site.objects.all()  # Assuming Site model exists
#     return render(request, 'emp/salary_preparation.html', {
#         'employees': employees,
#         'month': month,
#         'months': months,
#         'year': year,
#         'site': site,
#         'sites': sites,
#     })

@role_required('admin', 'hr')
def make_salary(request, employee_id):
    # print("request.POST.get('month')")
    # print("employee_id", employee_id)
    # print(request)
    employee = get_object_or_404(Employee, id=employee_id)
    month = int(request.POST.get('month', datetime.now().month))
    # month = 8 # For testing, hardcoded to August
    print('month', month)
    year = int(request.POST.get('year', datetime.now().year))

    # Check if already prepared
    if EmployeeSalaryTransaction.objects.filter(employee=employee, month=month, year=year).exists():
        return JsonResponse({'status': 'error', 'message': 'Salary already prepared for this month'})

    # Get active salary 2025-05-01
    active_salary = EmployeeSalaryMaster.objects.filter(
        employee=employee, status='Active', effective_from__lte=datetime(year, month, 1)
    ).order_by('-effective_from').first()

    if not active_salary:
        return JsonResponse({'status': 'error', 'message': 'No active salary master found'})

    salary_amount = active_salary.salary_amount

    # Get adjustments for the month
    adjustments = EmployeeAdjustment.objects.filter(
        employee=employee,
        date__month=month,
        date__year=year,
    ).annotate(
        signed_amount=Coalesce(
            models.Case(
                models.When(adjustment_type='Allowance', then=F('amount')),
                models.When(adjustment_type__in=['Deduction', 'Advance'], then=-F('amount')),
                output_field=DecimalField(),
            ), Value(Decimal('0'))
        )
    ).aggregate(net_adjust=Sum('signed_amount'))['net_adjust'] or Decimal('0')

    form = SalaryPreparationForm(request.POST if request.method == 'POST' else None)
    if request.method == 'POST' and form.is_valid():
        leave_days = form.cleaned_data['leave_days']
        remarks = form.cleaned_data['remarks']

        # Assume 30 days month
        per_day = salary_amount / Decimal('30')
        leave_deduction = per_day * Decimal(leave_days)

        net_salary = salary_amount + adjustments - leave_deduction

        transaction = EmployeeSalaryTransaction(
            employee=employee,
            month=month,
            year=year,
            salary_amount=salary_amount,
            adjustments_net=adjustments,
            leave_days=leave_days,
            leave_deduction=leave_deduction,
            net_salary=net_salary,
            prepared_by=request.user,
            remarks=remarks
        )
        transaction.save()
        return JsonResponse({'status': 'success', 'message': 'Salary prepared successfully'})

    # For GET, prepare data for modal
    previous_transactions = EmployeeSalaryTransaction.objects.filter(employee=employee).order_by('-year', '-month')[:3]
    current_adjustments = EmployeeAdjustment.objects.filter(employee=employee, date__month=month, date__year=year)
    
    data = {
        'salary_amount': float(salary_amount),
        'adjustments': list(current_adjustments.values()),
        'previous_transactions': list(previous_transactions.values()),
        'form': form.as_p(),  # Or use crispy
    }
    return JsonResponse(data)

# endregion SALARY_OLD

    
    
# region SALARY
@role_required('admin', 'hr')
def salary_master_list(request):
    # Employees without any salary
    no_salary_employees = Employee.objects.filter(salary_masters__isnull=True, status='Active')
    # Employees with latest active salary
    with_salary_employees = Employee.objects.filter(status='Active').annotate(
        latest_salary=Max('salary_masters__effective_from')
    ).filter(salary_masters__effective_from=F('latest_salary'), salary_masters__status='Active').prefetch_related('salary_masters').distinct()
    return render(request, 'emp/salary_master_list.html', {
        'no_salary_employees': no_salary_employees,
        'with_salary_employees': with_salary_employees,
        'view_name': request.resolver_match.view_name
    })

@role_required('admin', 'hr')
def salary_master_detail(request, emp_id):
    employee = get_object_or_404(Employee, id=emp_id)
    salaries = EmployeeSalaryMaster.objects.filter(employee=employee).order_by('-effective_from')
    latest_salary = salaries.first()
    if request.method == 'POST':
        if latest_salary and 'edit' in request.POST:
            form = EmployeeSalaryMasterForm(request.POST, instance=latest_salary)
        else:
            form = EmployeeSalaryMasterForm(request.POST)
        if form.is_valid():
            salary = form.save(commit=False)
            salary.employee = employee
            salary.created_by = request.user
            salary.save()
            return JsonResponse({
                'status': 'success',
                'message': 'Salary record saved successfully!'
            })
        else:
            return JsonResponse({
                'status': 'error', 
                'errors': form.errors,
                'message': 'Please correct the errors below.'
            }, status=400)
    else:
        form = EmployeeSalaryMasterForm(instance=latest_salary) if latest_salary else EmployeeSalaryMasterForm()
    return render(request, 'emp/salary_master_detail.html', {
        'employee': employee,
        'salaries': salaries,
        'form': form,
        'latest_salary': latest_salary,
        'view_name': request.resolver_match.view_name
    })

@role_required('admin', 'hr')
def adjustments_page(request):
    if request.method == 'POST':
        form = EmployeeAdjustmentForm(request.POST)
        if form.is_valid():
            adjustment = form.save(commit=False)
            adjustment.created_by = request.user
            adjustment.save()
            return redirect('apps.emp:adjustments_page')
    else:
        form = EmployeeAdjustmentForm()
    adjustments = EmployeeAdjustment.objects.all()
    return render(request, 'emp/adjustments.html', {'form': form, 'adjustments': adjustments, 'view_name': request.resolver_match.view_name})

@role_required('admin', 'hr')
def salary_preparation(request):
    from dateutil.relativedelta import relativedelta
    
    # Get selected month or default to previous month
    today = datetime.now()
    prev_month_date = today - relativedelta(months=1)
    
    selected_month = int(request.GET.get('month', prev_month_date.month))
    selected_year = int(request.GET.get('year', prev_month_date.year))
    
    # Generate 4 months (current + past 3)
    today = datetime.now()

    # Get Active Employee Count for status calculation
    # Only count employees active during that specific month ideally, but using current active as baseline is standard practice
    active_emp_count = Employee.objects.filter(status='Active').count()

    months_list = []
    for i in range(4):
        m_date = today - relativedelta(months=i)
        
        # Count transactions for this month
        trans_qs = EmployeeSalaryTransaction.objects.filter(
            month=m_date.month,
            year=m_date.year
        )
        total_prepared = trans_qs.count()
        total_processed = trans_qs.filter(status='Processed').count()
        
        # Determine Status
        status_code = 'Pending'
        if total_prepared == 0:
            status_code = 'Pending'
        elif total_processed == active_emp_count and active_emp_count > 0:
            status_code = 'Processed'
        elif total_prepared < active_emp_count:
            status_code = 'Partial'
        elif total_prepared >= active_emp_count: # All prepared, but not all processed
            status_code = 'Prepared'
            
        months_list.append({
            'month': m_date.month,
            'year': m_date.year,
            'display': m_date.strftime('%b %Y'),
            'status': status_code, # New field
            'prepared': total_prepared > 0, # Keep for backward compatibility if needed
            'is_selected': (m_date.month == selected_month and m_date.year == selected_year)
        })
    
    # Days in selected month
    days_in_month = calendar.monthrange(selected_year, selected_month)[1]
    
    # Filter employees
    site = request.GET.get('site')
    department = request.GET.get('department')
    designation = request.GET.get('designation')
    
    # Base query: Active employees with Active salary master
    queryset = Employee.objects.filter(
        status='Active',
        salary_masters__status='Active'
    ).distinct()
    
    # Apply filters
    if site:
        queryset = queryset.filter(site__id=site)
    if department:
        queryset = queryset.filter(department__id=department)
    if designation:
        queryset = queryset.filter(designation__id=designation)
    
    # Annotate with adjustment data using NEW models
    queryset = queryset.annotate(
        # Cleared adjustments IN the selected month
        cleared_adjustments=Coalesce(
            Sum('employeeadjustmenttransaction__settlement_amount',
                filter=Q(
                    employeeadjustmenttransaction__salary_month=selected_month,
                    employeeadjustmenttransaction__salary_year=selected_year
                )),
            Decimal('0')
        )
        # Note: pending_adjustments removed from annotation to prevent duplication
        # Will be calculated in Python loop below
    )
    
    # Add salary transaction info for each employee
    for emp in queryset:
        # Last salary transaction (any month)
        emp.last_salary_obj = EmployeeSalaryTransaction.objects.filter(
            employee=emp
        ).order_by('-year', '-month').first()
        
        # Selected month salary transaction
        emp.selected_month_trans = EmployeeSalaryTransaction.objects.filter(
            employee=emp,
            month=selected_month,
            year=selected_year
        ).first()
        
        # Current month info
        emp.is_prepared = bool(emp.selected_month_trans)
        if emp.selected_month_trans:
            emp.current_month_net = emp.selected_month_trans.net_salary
            emp.selected_month_absent = emp.selected_month_trans.leave_days
            # Can edit only if status is 'Prepared' (not yet processed)
            emp.can_edit = emp.selected_month_trans.status == 'Prepared'
        else:
            emp.current_month_net = None
            emp.selected_month_absent = 0
            emp.can_edit = False
        
        # Calculate pending adjustments correctly (avoid double-counting from joins)
        active_masters = EmployeeAdjustmentMaster.objects.filter(
            employee=emp,
            status='Active'
        )
        emp.pending_adjustments = sum(
            master.outstanding_amount for master in active_masters
        )
    
    sites = Site.objects.all()
    departments = Department.objects.all()
    designations = Designation.objects.all()
    
    # Calculate statistics for report header
    total_employees = queryset.count()
    prepared_count = sum(1 for emp in queryset if emp.is_prepared)
    not_prepared_count = total_employees - prepared_count
    processed_count = sum(1 for emp in queryset if emp.is_prepared and not emp.can_edit)
    
    # Get selected filter labels
    selected_site_name = Site.objects.get(id=site).site_name if site else 'All Sites'
    selected_dept_name = Department.objects.get(id=department).department_name if department else 'All Departments'
    selected_desig_name = Designation.objects.get(id=designation).designation_name if designation else 'All Designations'
    
    return render(request, 'emp/salary_preparation.html', {
        'months_list': months_list,
        'selected_month': selected_month,
        'selected_year': selected_year,
        'selected_month_name': calendar.month_name[selected_month],
        'days_in_month': days_in_month,
        'employees': queryset,
        'sites': sites,
        'departments': departments,
        'designations': designations,
        'view_name': request.resolver_match.view_name,
        # Statistics for report header
        'total_employees': total_employees,
        'prepared_count': prepared_count,
        'not_prepared_count': not_prepared_count,
        'processed_count': processed_count,
        'selected_site_name': selected_site_name,
        'selected_dept_name': selected_dept_name,
        'selected_desig_name': selected_desig_name,
    })





@role_required('admin', 'hr')
def process_salary_bulk(request):
    """
    Bulk process selected employee salaries.
    Updates status from 'Prepared' to 'Processed' for selected employees.
    """
    
    
    if request.method == 'POST':
        employee_ids = request.POST.getlist('employee_ids')
        month = int(request.POST.get('month'))
        year = int(request.POST.get('year'))
        
        if not employee_ids:
            messages.error(request, 'No employees selected.')
            return redirect('emp:salary_preparation')
        
        # Update salary transactions to Processed
        updated_count = EmployeeSalaryTransaction.objects.filter(
            employee_id__in=employee_ids,
            month=month,
            year=year,
            status='Prepared'  # Only update Prepared ones
        ).update(status='Processed')
        
        if updated_count > 0:
            messages.success(request, f'Successfully processed {updated_count} salary record(s).')
        else:
            messages.warning(request, 'No salaries were processed. They may already be processed.')
        return redirect(f'/emp/salary-preparation/?month={month}&year={year}')
    
    return redirect('emp:salary_preparation')


@role_required('admin', 'hr')
def export_salary_xlsx(request):
    """Export salary preparation report to Excel"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        messages.error(request, 'openpyxl library not installed.')
        return redirect('emp:salary_preparation')
    
    month = int(request.GET.get('month', datetime.now().month))
    year = int(request.GET.get('year', datetime.now().year))
    site = request.GET.get('site', '')
    department = request.GET.get('department', '')
    designation = request.GET.get('designation', '')
    
    queryset = Employee.objects.filter(status='Active', salary_masters__status='Active').distinct()
    if site:
        queryset = queryset.filter(site__id=site)
    if department:
        queryset = queryset.filter(department__id=department)
    if designation:
        queryset = queryset.filter(designation__id=designation)
    
    queryset = queryset.annotate(
        cleared_adjustments=Coalesce(
            Sum('employeeadjustmenttransaction__settlement_amount',
                filter=Q(employeeadjustmenttransaction__salary_month=month, employeeadjustmenttransaction__salary_year=year)),
            Decimal('0'))
    )
    
    for emp in queryset:
        emp.last_salary_obj = EmployeeSalaryTransaction.objects.filter(employee=emp).order_by('-year', '-month').first()
        emp.selected_month_trans = EmployeeSalaryTransaction.objects.filter(employee=emp, month=month, year=year).first()
        emp.is_prepared = bool(emp.selected_month_trans)
        if emp.selected_month_trans:
            emp.current_month_net = emp.selected_month_trans.net_salary
            emp.can_edit = emp.selected_month_trans.status == 'Prepared'
        else:
            emp.current_month_net = None
            emp.can_edit = False
        active_masters = EmployeeAdjustmentMaster.objects.filter(employee=emp, status='Active')
        emp.pending_adjustments = sum(master.outstanding_amount for master in active_masters)
    
    total_employees = queryset.count()
    prepared_count = sum(1 for emp in queryset if emp.is_prepared)
    selected_site_name = Site.objects.get(id=site).site_name if site else 'All Sites'
    selected_dept_name = Department.objects.get(id=department).department_name if department else 'All Departments'
    selected_desig_name = Designation.objects.get(id=designation).designation_name if designation else 'All Designations'
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Salary Report"
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    title_font = Font(bold=True, size=14)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:I1')
    ws['A1'] = 'Salary Preparation Report'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center')
    
    row = 3
    ws[f'A{row}'] = f'Period: {calendar.month_name[month]} {year}'
    ws[f'A{row}'].font = Font(bold=True)
    row += 1
    ws[f'A{row}'] = f'Site: {selected_site_name}'
    ws[f'D{row}'] = f'Department: {selected_dept_name}'
    ws[f'G{row}'] = f'Designation: {selected_desig_name}'
    row += 2
    
    headers = ['S.No', 'Employee', 'Department', 'Designation', 'Current Salary', 'Cleared', 'Pending', 'Last Salary', 'Status']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center')
    
    row += 1
    for idx, emp in enumerate(queryset, 1):
        ws.cell(row=row, column=1, value=idx).border = border
        ws.cell(row=row, column=2, value=f"{emp.first_name} {emp.last_name}").border = border
        ws.cell(row=row, column=3, value=emp.department.department_name if emp.department else '').border = border
        ws.cell(row=row, column=4, value=emp.designation.designation_name if emp.designation else '').border = border
        ws.cell(row=row, column=5, value=float(emp.current_month_net) if emp.current_month_net else 0).border = border
        ws.cell(row=row, column=6, value=float(emp.cleared_adjustments)).border = border
        ws.cell(row=row, column=7, value=float(emp.pending_adjustments)).border = border
        ws.cell(row=row, column=8, value=float(emp.last_salary_obj.net_salary) if emp.last_salary_obj else 0).border = border
        status = 'Processed' if (emp.is_prepared and not emp.can_edit) else ('Prepared' if emp.is_prepared else 'Not Prepared')
        ws.cell(row=row, column=9, value=status).border = border
        row += 1
    
    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
        ws.column_dimensions[col].width = 20 if col == 'B' else 15
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=Salary_Report_{month}_{year}.xlsx'
    wb.save(response)
    return response


@role_required('admin', 'hr')
def export_salary_pdf(request):
    """Export salary preparation report to PDF"""
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
    except ImportError:
        messages.error(request, 'reportlab library not installed.')
        return redirect('emp:salary_preparation')
    
    month = int(request.GET.get('month', datetime.now().month))
    year = int(request.GET.get('year', datetime.now().year))
    site = request.GET.get('site', '')
    department = request.GET.get('department', '')
    designation = request.GET.get('designation', '')
    
    queryset = Employee.objects.filter(status='Active', salary_masters__status='Active').distinct()
    if site:
        queryset = queryset.filter(site__id=site)
    if department:
        queryset = queryset.filter(department__id=department)
    if designation:
        queryset = queryset.filter(designation__id=designation)
    
    queryset = queryset.annotate(
        cleared_adjustments=Coalesce(
            Sum('employeeadjustmenttransaction__settlement_amount',
                filter=Q(employeeadjustmenttransaction__salary_month=month, employeeadjustmenttransaction__salary_year=year)),
            Decimal('0'))
    )
    
    for emp in queryset:
        emp.last_salary_obj = EmployeeSalaryTransaction.objects.filter(employee=emp).order_by('-year', '-month').first()
        emp.selected_month_trans = EmployeeSalaryTransaction.objects.filter(employee=emp, month=month, year=year).first()
        emp.is_prepared = bool(emp.selected_month_trans)
        if emp.selected_month_trans:
            emp.current_month_net = emp.selected_month_trans.net_salary
            emp.can_edit = emp.selected_month_trans.status == 'Prepared'
        else:
            emp.current_month_net = None
            emp.can_edit = False
        active_masters = EmployeeAdjustmentMaster.objects.filter(employee=emp, status='Active')
        emp.pending_adjustments = sum(master.outstanding_amount for master in active_masters)
    
    selected_site_name = Site.objects.get(id=site).site_name if site else 'All Sites'
    selected_dept_name = Department.objects.get(id=department).department_name if department else 'All Departments'
    selected_desig_name = Designation.objects.get(id=designation).designation_name if designation else 'All Designations'
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename=Salary_Report_{month}_{year}.pdf'
    
    doc = SimpleDocTemplate(response, pagesize=landscape(A4))
    elements = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], alignment=1)
    elements.append(Paragraph('Salary Preparation Report', title_style))
    elements.append(Spacer(1, 0.3 * inch))
    
    header_data = [[f"Period: {calendar.month_name[month]} {year}", '', ''], [f"Site: {selected_site_name}", f"Dept: {selected_dept_name}", f"Desig: {selected_desig_name}"]]
    header_table = Table(header_data, colWidths=[3*inch, 3*inch, 3*inch])
    header_table.setStyle(TableStyle([('FONT', (0, 0), (-1, -1), 'Helvetica-Bold', 10), ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)]))
    elements.append(header_table)
    elements.append(Spacer(1, 0.3 * inch))
    
    data = [['S.No', 'Employee', 'Dept', 'Desig', 'Current', 'Cleared', 'Pending', 'Last', 'Status']]
    for idx, emp in enumerate(queryset, 1):
        status = 'Proc' if (emp.is_prepared and not emp.can_edit) else ('Prep' if emp.is_prepared else 'Not')
        data.append([idx, f"{emp.first_name} {emp.last_name}"[:20], emp.department.department_name[:12] if emp.department else '', 
                     emp.designation.designation_name[:12] if emp.designation else '', f"₹{emp.current_month_net:,.0f}" if emp.current_month_net else '₹0',
                     f"₹{emp.cleared_adjustments:,.0f}", f"₹{emp.pending_adjustments:,.0f}",
                     f"₹{emp.last_salary_obj.net_salary:,.0f}" if emp.last_salary_obj else '₹0', status])
    
    table = Table(data, colWidths=[0.5*inch, 1.5*inch, 1*inch, 1*inch, 1*inch, 1*inch, 1*inch, 1*inch, 0.7*inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
    ]))
    
    elements.append(table)
    doc.build(elements)
    return response


@role_required('admin', 'hr')
def prepare_salary(request, employee_id):
    from apps.emp.models import EmployeeAdjustmentMaster, EmployeeAdjustmentTransaction
    
    employee = get_object_or_404(Employee, id=employee_id)
    
    # Get month/year from query params or default to current
    selected_month = int(request.GET.get('month', datetime.now().month))
    selected_year = int(request.GET.get('year', datetime.now().year))
    
    days_in_month = calendar.monthrange(selected_year, selected_month)[1]
    latest_salary = EmployeeSalaryMaster.objects.filter(employee=employee, status='Active').order_by('-effective_from').first()
    if not latest_salary:
        return JsonResponse({'status': 'error', 'message': 'No active salary'}, status=400)
    
    # Load ACTIVE adjustment masters (with outstanding balance)
    pending_adjustments = EmployeeAdjustmentMaster.objects.filter(
        employee=employee, 
        status='Active'
    )
    
    # Get last 3 salaries for context
    last_3_salaries = EmployeeSalaryTransaction.objects.filter(employee=employee).order_by('-year', '-month')[:3]
    
    if request.method == 'POST':
        # Get month/year from form submission
        month = int(request.POST.get('month', selected_month))
        year = int(request.POST.get('year', selected_year))
        days_in_month_calc = calendar.monthrange(year, month)[1]
        
        days_absent = int(request.POST.get('days_absent', 0))
        leave_deduction = (latest_salary.salary_amount / days_in_month_calc) * days_absent
        adjustments_net = Decimal('0')
        
        # Check if salary already prepared for this month
        existing = EmployeeSalaryTransaction.objects.filter(
            employee=employee,
            month=month,
            year=year
        ).first()
        
        if existing:
            return JsonResponse({
                'status': 'error',
                'message': f'Salary already prepared for {month}/{year}'
            }, status=400)
        
        # Process each pending adjustment master
        adjustment_details = []
        
        for master in pending_adjustments:
            settlement_amount = Decimal('0')
            outstanding = master.outstanding_amount
            
            # Check if full settlement selected
            if request.POST.get(f'settle_full_{master.id}'):
                settlement_amount = outstanding
                transaction_type = 'Full'
                adjustment_details.append(f"{master.adjustment_type} #{master.id}: Full ₹{settlement_amount}")
            
            # Check if partial settlement selected
            elif request.POST.get(f'partial_{master.id}'):
                partial_value = request.POST.get(f'partial_{master.id}', '0')
                try:
                    partial = Decimal(partial_value)
                except:
                    partial = Decimal('0')
                
                # Validation
                if partial > outstanding:
                    return JsonResponse({
                        'status': 'error',
                        'message': f'Partial amount (₹{partial}) exceeds outstanding (₹{outstanding}) for {master.adjustment_type}'
                    }, status=400)
                
                if partial > 0:
                    settlement_amount = partial
                    transaction_type = 'Partial'
                    adjustment_details.append(f"{master.adjustment_type} #{master.id}: Partial ₹{settlement_amount}")
            
            # Create settlement transaction if amount > 0
            if settlement_amount > 0:
                EmployeeAdjustmentTransaction.objects.create(
                    adjustment_master=master,
                    employee=employee,
                    settlement_amount=settlement_amount,
                    salary_month=month,
                    salary_year=year,
                    transaction_type=transaction_type,
                    remarks=f"Settled via {calendar.month_name[month]} {year} salary",
                    created_by=request.user
                )
                
                # Update cleared_amount in master
                master.cleared_amount += settlement_amount
                master.save()
                
                # Update master status if fully settled
                if master.is_fully_settled():
                    master.status = 'Closed'
                    master.save()
                
                # Calculate net adjustment (+ for Bonus/Reward, - for Advance/Fine)
                sign = 1 if master.adjustment_type in ['Bonus', 'Reward'] else -1
                adjustments_net += sign * settlement_amount
        
        # Calculate final salary
        net_salary = latest_salary.salary_amount + adjustments_net - leave_deduction
        
        # Build comprehensive transaction remarks
        transaction_remarks = f"Salary for {calendar.month_name[month]} {year}\n"
        transaction_remarks += f"Base Salary: ₹{latest_salary.salary_amount}\n"
        transaction_remarks += f"Leave Days: {days_absent}, Deduction: ₹{leave_deduction}\n"
        
        if adjustment_details:
            transaction_remarks += f"\nAdjustments Processed:\n"
            for detail in adjustment_details:
                transaction_remarks += f"- {detail}\n"
            transaction_remarks += f"Total Adjustments: ₹{adjustments_net}\n"
        else:
            transaction_remarks += "No adjustments for this month\n"
        
        transaction_remarks += f"\nNet Salary: ₹{net_salary}"
        
        # Create salary transaction
        transaction = EmployeeSalaryTransaction(
            employee=employee,
            month=month,
            year=year,
            salary_amount=latest_salary.salary_amount,
            adjustments_net=adjustments_net,
            leave_days=days_absent,
            leave_deduction=leave_deduction,
            net_salary=net_salary,
            prepared_by=request.user,
            remarks=transaction_remarks
        )
        transaction.save()
        return JsonResponse({'status': 'success', 'message': 'Salary prepared successfully'})
    
    # GET Request: Fetch existing transaction context
    existing_transaction = EmployeeSalaryTransaction.objects.filter(
        employee=employee,
        month=selected_month,
        year=selected_year
    ).first()
    
    show_salary_slip = False
    salary_slip_data = {}
    default_absent = 0
    
    if existing_transaction:
        if existing_transaction.status == 'Processed':
            show_salary_slip = True
            # Fetch adjustments settled in THIS salary
            adjustments_settled = EmployeeAdjustmentTransaction.objects.filter(
                employee=employee,
                salary_month=selected_month,
                salary_year=selected_year
            )
            salary_slip_data = {
                'transaction': existing_transaction,
                'adjustments': adjustments_settled,
                'earnings': latest_salary.salary_amount,
                'deductions': existing_transaction.leave_deduction,
                'net_pay': existing_transaction.net_salary
            }
        else:
            # Prepared but not Processed - Edit Mode
            default_absent = existing_transaction.leave_days
            
    return render(request, 'emp/prepare_salary_modal.html', {
        'employee': employee,
        'latest_salary': latest_salary,
        'pending_adjustments': pending_adjustments,  # Now contains Masters
        'last_3_salaries': last_3_salaries,
        'days_in_month': days_in_month,
        'selected_month': selected_month,
        'selected_year': selected_year,
        'selected_month_name': calendar.month_name[selected_month],
        'view_name': request.resolver_match.view_name,
        'default_absent': default_absent,
        'existing_transaction': existing_transaction,
        'show_salary_slip': show_salary_slip,
        'salary_slip_data': salary_slip_data
    })



@role_required('admin', 'hr')
def employee_profile_popup(request, employee_id):
    """Employee profile popup for HRMS dashboard"""
    from dateutil.relativedelta import relativedelta
    
    employee = get_object_or_404(Employee, id=employee_id)
    
    # Calculate duration since joining
    today = datetime.now().date()
    delta = relativedelta(today, employee.joining_date)
    if delta.years > 0:
        duration_str = f"{delta.years} yr {delta.months} mos"
    else:
        duration_str = f"{delta.months} mos"
    
    # Status checks
    has_login = User.objects.filter(username=employee.emp_code).exists()
    has_salary = EmployeeSalaryMaster.objects.filter(
        employee=employee, 
        status='Active'
    ).exists()
    
    # Payroll status for current month
    current_month = datetime.now().month
    current_year = datetime.now().year
    payroll_processed = EmployeeSalaryTransaction.objects.filter(
        employee=employee,
        month=current_month,
        year=current_year
    ).exists()
    
    # Last salary transaction
    last_salary = EmployeeSalaryTransaction.objects.filter(
        employee=employee
    ).order_by('-year', '-month').first()
    
    # Get all adjustment masters for this employee (both active and closed)
    from apps.emp.models import EmployeeAdjustmentMaster, EmployeeAdjustmentTransaction
    
    adjustment_masters = EmployeeAdjustmentMaster.objects.filter(
        employee=employee
    ).order_by('-date_issued')[:5]  # Last 5 adjustments
    
    # Get ALL transactions for display (last 10)
    all_transactions = EmployeeAdjustmentTransaction.objects.filter(
        employee=employee
    ).select_related('adjustment_master').order_by('-salary_year', '-salary_month', '-settlement_date')[:10]
    
    context = {
        'employee': employee,
        'duration_str': duration_str,
        'has_login': has_login,
        'has_salary': has_salary,
        'payroll_processed': payroll_processed,
        'last_salary': last_salary,
        'adjustment_masters': adjustment_masters,  # NEW: Masters with outstanding info
        'all_transactions': all_transactions,  # NEW: Transaction history
        'view_name': request.resolver_match.view_name
    }
    
    return render(request, 'emp/employee_profile_popup.html', context)

# endregion SALARY

@role_required('admin', 'hr')
def get_employees_json(request):
    """API endpoint for Select2 employee search"""
    q = request.GET.get('q', '')
    page = int(request.GET.get('page', 1))
    page_size = 20
    
    employees = Employee.objects.filter(status='Active')
    
    if q:
        employees = employees.filter(
            Q(first_name__icontains=q) | 
            Q(last_name__icontains=q) | 
            Q(emp_code__icontains=q)
        )
    
    total = employees.count()
    start = (page - 1) * page_size
    end = start + page_size
    
    employees_page = employees[start:end]
    
    results = [{
        'id': emp.id,
        'text': f"{emp.first_name} {emp.last_name}",
        'emp_code': emp.emp_code
    } for emp in employees_page]
    
    return JsonResponse({
        'results': results,
        'pagination': {
            'more': end < total
        }
    })

@role_required('admin', 'hr')
def get_employee_profile(request, employee_id):
    """API endpoint to get employee profile details"""
    try:
        employee = Employee.objects.get(id=employee_id, status='Active')
        
        # Check if user exists
        try:
            user = User.objects.get(username=employee.emp_code)
            has_user = True
            username = user.username
            user_id = user.id
            user_role = user.groups.first().name if user.groups.exists() else 'Employee'
        except User.DoesNotExist:
            has_user = False
            username = None
            user_id = None
            user_role = None
        
        data = {
            'name': f"{employee.first_name} {employee.last_name}",
            'emp_code': employee.emp_code,
            'designation': employee.designation.designation_name,
            'department': employee.department.department_name,
            'email': employee.email or 'N/A',
            'contact_no': employee.contact_no,
            'joining_date': employee.joining_date.strftime('%Y-%m-%d'),
            'photo': employee.photo.url if employee.photo else '/static/dist/img/default-150x150.png',
            'has_user': has_user,
            'username': username,
            'user_id': user_id,
            'user_role': user_role
        }
        
        return JsonResponse(data)
    except Employee.DoesNotExist:
        return JsonResponse({'error': 'Employee not found'}, status=404)


@role_required('admin', 'hr')
def get_employees_list(request):
    """API endpoint to get all employees with login status for create user page"""
    employees = Employee.objects.filter(status='Active').order_by('first_name', 'last_name')
    
    employee_data = []
    for emp in employees:
        # Check if user exists
        try:
            user = User.objects.get(username=emp.emp_code)
            has_login = True
        except User.DoesNotExist:
            has_login = False
        
        employee_data.append({
            'id': emp.id,
            'name': f"{emp.first_name} {emp.last_name}",
            'emp_code': emp.emp_code,
            'designation': emp.designation.designation_name,
            'department': emp.department.department_name,
            'photo': emp.photo.url if emp.photo else None,
            'has_login': has_login
        })
    
    # Sort: employees without login first
    employee_data.sort(key=lambda x: (x['has_login'], x['name']))
    
    return JsonResponse({'employees': employee_data})

@role_required('admin', 'hr')
def check_duplicate(request):
    """API endpoint to check for duplicate Aadhaar/PAN/Mobile"""
    field = request.GET.get('field')  # 'aadhaar_no', 'pan_card', 'contact_no'
    value = request.GET.get('value')
    employee_id = request.GET.get('employee_id')  # For edit mode
    
    if not field or not value:
        return JsonResponse({'exists': False})
    
    # Build query
    query_filter = {field: value}
    
    # Check if exists (excluding current employee if editing)
    queryset = Employee.objects.filter(**query_filter)
    if employee_id:
        queryset = queryset.exclude(id=employee_id)
    
    exists = queryset.exists()
    
    return JsonResponse({'exists': exists})


from django.core.paginator import Paginator
from dateutil.relativedelta import relativedelta

def employee_profile_popup(request, employee_id):
    employee = get_object_or_404(Employee, id=employee_id)
    
    # Check User Login
    try:
        User.objects.get(username=employee.emp_code)
        has_login = True
    except User.DoesNotExist:
        has_login = False
        
    # Check Salary Master
    has_salary = EmployeeSalaryMaster.objects.filter(employee=employee, status='Active').exists()
    
    # Check Last Salary Processed
    last_salary = EmployeeSalaryTransaction.objects.filter(employee=employee).order_by('-year', '-month').first()
    payroll_processed = last_salary is not None
    
    # Duration Calculation
    today = date.today()
    if employee.joining_date:
        duration = relativedelta(today, employee.joining_date)
        duration_str = f"{duration.years}y {duration.months}m"
    else:
        duration_str = "N/A"

    # Salaries History (Paginated)
    salaries_list = EmployeeSalaryTransaction.objects.filter(employee=employee).order_by('-year', '-month')
    paginator_sal = Paginator(salaries_list, 5)
    page_sal = request.GET.get('page_sal')
    salaries = paginator_sal.get_page(page_sal)
    
    # Adjustments (Paginated)
    adj_list = EmployeeAdjustmentMaster.objects.filter(employee=employee).order_by('-created_at')
    paginator_adj = Paginator(adj_list, 5)
    page_adj = request.GET.get('page_adj')
    adjustments = paginator_adj.get_page(page_adj)
    
    context = {
        'employee': employee,
        'has_login': has_login,
        'has_salary': has_salary,
        'payroll_processed': payroll_processed,
        'salaries': salaries,
        'adjustments': adjustments,
        'duration_str': duration_str,
    }
    return render(request, 'emp/employee_profile_popup.html', context)

@role_required('admin')
def admin_role_assignment(request):
    """
    Admin page to manage user roles
    """
    if request.method == 'POST':
        user_id = request.POST.get('user_id')
        new_role = request.POST.get('new_role')
        
        try:
            target_user = User.objects.get(id=user_id)
            # Don't allow changing superuser roles here safely
            if target_user.is_superuser:
                messages.error(request, "Cannot modify Superuser roles from this panel.")
            else:
                # Clear all existing groups
                target_user.groups.clear()
                
                # Add new group
                if new_role in ['admin', 'hr', 'manager', 'employee']:
                    try:
                        group = Group.objects.get(name=new_role)
                        target_user.groups.add(group)
                        messages.success(request, f"Role updated to '{new_role}' for {target_user.username}")
                    except Group.DoesNotExist:
                         messages.error(request, f"Role group '{new_role}' does not exist.")
                else:
                    messages.error(request, "Invalid role selected.")
                    
        except User.DoesNotExist:
            messages.error(request, "User not found.")
            
        return redirect('emp:admin_role_assignment')

    # GET Request
    # Fetch all users who are not superusers
    users_qs = User.objects.filter(is_superuser=False).order_by('username')
    
    users_data = []
    for u in users_qs:
        # Get current role (first group found)
        current_role = u.groups.first().name if u.groups.exists() else 'No Role'
        
        # Try to find employee details if username matches emp_code
        emp_name = "N/A"
        try:
            emp = Employee.objects.get(emp_code=u.username)
            emp_name = f"{emp.first_name} {emp.last_name}"
        except Employee.DoesNotExist:
            pass
            
        users_data.append({
            'id': u.id,
            'username': u.username,
            'email': u.email,
            'name': emp_name,
            'role': current_role
        })

    context = {
        'users': users_data,
        'all_roles': ['admin', 'hr', 'manager', 'employee']
    }
    return render(request, 'emp/admin_role_assignment.html', context)
